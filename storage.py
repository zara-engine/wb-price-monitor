"""Хранение состояния, поиск изменений и Excel-отчёт."""

from __future__ import annotations

import json
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from wb_client import Product

HEADERS = ["Дата", "Артикул", "Название", "Бренд", "Продавец", "Цена, ₽",
           "Старая цена, ₽", "Скидка, %", "Остаток", "Рейтинг", "Отзывы", "Ссылка"]
WIDTHS = [17, 13, 48, 14, 20, 15, 18, 12, 11, 11, 11, 55]


@dataclass
class Change:
    product: Product
    kind: str             # price | out | back
    old_price: float | None = None

    @property
    def percent(self) -> float | None:
        if self.kind != "price" or not self.old_price or self.product.price is None:
            return None
        return (self.product.price - self.old_price) / self.old_price * 100


def load_state(path: Path) -> dict:
    return json.loads(path.read_text("utf-8")) if path.exists() else {}


def save_state(path: Path, products: list[Product]) -> None:
    state = {str(p.nm): {"price": p.price, "in_stock": p.in_stock} for p in products}
    path.write_text(json.dumps(state, ensure_ascii=False, indent=2), "utf-8")


def find_changes(products: list[Product], state: dict,
                 threshold: float = 0.0) -> list[Change]:
    """Сравнивает новые данные с прошлым запуском.

    threshold — минимальное изменение цены в процентах, о котором сообщать.
    """
    changes: list[Change] = []
    for p in products:
        old = state.get(str(p.nm))
        if old is None:
            continue  # первый запуск для этого артикула
        if old["in_stock"] and not p.in_stock:
            changes.append(Change(p, "out"))
        elif not old["in_stock"] and p.in_stock:
            changes.append(Change(p, "back"))
        elif p.in_stock and old["price"] and p.price != old["price"]:
            ch = Change(p, "price", old["price"])
            if abs(ch.percent or 0) >= threshold:
                changes.append(ch)
    return changes


def history_rows(products: list[Product]) -> list[list]:
    """Строки истории по колонкам HEADERS (общие для Excel и Google Таблиц)."""
    now = datetime.now().strftime("%d.%m.%Y %H:%M")
    rows = []
    for p in products:
        price = p.price if p.in_stock else "нет в наличии"
        rows.append([now, p.nm, p.name, p.brand, p.supplier, price,
                     p.basic_price if p.in_stock else None,
                     p.discount if p.in_stock else None,
                     p.quantity, p.rating, p.feedbacks, p.url])
    return rows


def append_history(path: Path, products: list[Product]) -> None:
    """Дописывает строки в Excel-историю (создаёт файл при первом запуске)."""
    if path.exists():
        wb = load_workbook(path)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "История цен"
        ws.append(HEADERS)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="7B2CBF")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for i, w in enumerate(WIDTHS, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"

    for row in history_rows(products):
        ws.append(row)
        ws.cell(ws.max_row, 12).hyperlink = row[-1]
    wb.save(path)
